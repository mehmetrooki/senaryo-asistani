"""Fransız senaryo formatında örnek .docx dosyaları üretir (deneme için)."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

HERE = Path(__file__).parent


def build(lines, path):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Courier New"
    st.font.size = Pt(12)
    for kind, text in lines:
        p = doc.add_paragraph()
        r = p.add_run(text)
        if kind == "sahne":
            r.bold = True
        elif kind == "karakter":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.left_indent = Cm(0)
        elif kind == "diyalog":
            p.paragraph_format.left_indent = Cm(3)
            p.paragraph_format.right_indent = Cm(3)
        elif kind == "aksiyon":
            r.italic = True
    doc.save(path)


V1 = [
    ("sahne", "1. İÇ. LABORATUVAR - GÜN"),
    ("aksiyon", "Robo ve Mina kocaman bir makinenin önünde duruyor. Alarm çalıyor."),
    ("karakter", "MİNA"),
    ("diyalog", "Ha, bunu gör! Makine kudurmuş gibi titriyor."),
    ("karakter", "ROBO"),
    ("diyalog", "Hemen kırmızı kabloyu kes, yoksa her şey patlayacak!"),
    ("aksiyon", "Mina makasla kabloya uzanır."),
    ("karakter", "MİNA"),
    ("diyalog", "Sen tam bir aptalsın Robo, ben hallederim."),
    ("sahne", "2. DIŞ. BAHÇE - AKŞAM"),
    ("aksiyon", "İkisi de yorgun ama mutlu, çimenlere oturur."),
]

V2 = [
    ("sahne", "1. İÇ. LABORATUVAR - GÜN"),
    ("aksiyon", "Robo ve Mina kocaman bir makinenin önünde duruyor. Alarm çalıyor."),
    ("karakter", "MİNA"),
    ("diyalog", "Bak, bunu gör! Makine sallanıp duruyor."),
    ("karakter", "ROBO"),
    ("diyalog", "Hemen kırmızı butona bas, yoksa her şey duracak!"),
    ("aksiyon", "Mina butona uzanır."),
    ("karakter", "MİNA"),
    ("diyalog", "Sen tam bir şaşkınsın Robo, ben hallederim."),
    ("sahne", "2. DIŞ. BAHÇE - AKŞAM"),
    ("aksiyon", "İkisi de yorgun ama mutlu, çimenlere oturur."),
]

V3 = [
    ("sahne", "5. İÇ. MUTFAK - SABAH"),
    ("aksiyon", "Tosun tezgâhın üstüne tırmanmış, prize çatal sokmaya çalışıyor."),
    ("karakter", "TOSUN"),
    ("diyalog", "Ha! Bakın ne yapıyorum!"),
    ("karakter", "MİNA"),
    ("diyalog", "Çeneni kapa da in oradan, kudurdun mu sen?"),
]

if __name__ == "__main__":
    build(V1, HERE / "Bolum12_v1.docx")
    build(V2, HERE / "Bolum12_v2.docx")
    build(V3, HERE / "Bolum13.docx")
    print("örnekler hazır")
